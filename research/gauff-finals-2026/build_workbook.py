from pathlib import Path
import sys,json,hashlib
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
R=Path('/workspace/gauff-finals-production');REPO=Path('/workspace/tennislive-gauff');sys.path.insert(0,str(REPO/'src'));from tennislive.zh import player_zh
D=json.loads((R/'facts/canonical.json').read_text());O=json.loads((R/'editorial/outline.json').read_text());wb=Workbook();s=wb.active;s.title='口径与核验'
notes=[['项目','说明'],['球员与赛季','Coco Gauff 高芙／2026'],['统计截止',D['generated_at']],['实打与行政记录','16站；62场实打47胜15负；11BYE、1对手W/O独立不计实打；IW本人RET为正常计战绩的一负。'],['冠军与决赛','1冠：辛辛那提。2次亚军：迈阿密、罗马；共3次巡回赛单打决赛。'],['三盘战绩','18胜11负，29场正常打满三盘的单打；RET/BYE/W/O不混入三盘统计。'],['积分','16站各只计一次＝5774；10/19官方Race第5/5774一致。逐场每行携带的本站积分不是单场所得。'],['比赛日口径','WTA赛事UTC开赛时间转赛地当地日与北京时间；不将LastUpdated、报道时间或开始+比赛用时称作实际末分时刻。'],['联合杯','官方WTA2084赛事feed为空；4场当地比赛日由UnitedCup官方逐日赛报/结果页证实，精确UTC开赛未知不编造。单打3胜1负，美国队四强，非个人单打冠军。'],['Toronto W/O','1/4决赛本西奇赛前退赛，高芙晋级但不计入实打胜负，不伪造该场比分或赛点。'],['罗马抢七冲突','历史接口与赛事ScoreTbSet2均输家3分；官方赛报4504616写7-5。原始数据保留，待实际视频解决；展示比分先省抢七小分。'],['温网抢七字段','历史6-7(10)，赛事ScoreTbSet3=10，而ScoreString为7-6(12)。括号记录需分清输家/赢家；展示前还核真实画面。'],['姓名ID差异','Wimbledon Claire Liu：history323966，赛事feed179300；由官方全名Claire Liu精确对应，非模糊拼接。'],['年终确认','WTA官方10/14 10:00Z＝北京时间18:00公告，连续第5年获资格。CMS publishFrom10/13不是最早公开确认或数学锁定证据。'],['美网以后','10/7北京第四轮负梅尔滕斯；随后已宣布退出武汉。本人称漫长一年后需要休息恢复，不编造具体伤处。'],['未来边界','没有把年终或其他未来未赛赛事写成已发生结果；不把最后一场可见赛果等同赛季已结束。'],['制作重点','四GS＋辛辛那提冠军＋迈阿密/罗马RU共7站；其余9站简短旁白。片长约5min但完整全季数据仍保留。'],['原始逐场来源',D['matches'][0]['player_api_source']]]
for r in notes:s.append(r)
def cn(name):
 try:return player_zh(name)
 except Exception:return name
s=wb.create_sheet('16站汇总');s.append(['赛事','赛事起日','赛事止日','级别API','场地','成绩','实打胜','实打负','BYE','RET','W/O','赛事最终积分只计一次','最后对手','最后比分高芙视角','末场赛地当地日','末场北京时间日','重点或略提','官方日期来源'])
focus={e['event_id']for e in O['focus_events']}
for e in D['events']:
 s.append([e['event'],e['start'],e['end'],e['level'],e['surface'],e['finish'],e['wins'],e['losses'],e['bye'],e['ret'],e['wo'],e['points'],cn(e['last_opponent']),e['last_score_player'],e['final_date'],e['final_date_bjt'],'重点'if e['event_id']in focus else'略提',e['date_source']])
heads=['赛事','轮次','代码','赛地当地日','北京时间日','UTC日','日期证据方法','UTC开赛或预计时间','当地开赛时间','北京时间开赛时间','精确比赛结束时间','赛地时区','对手中文','对手英文','对手当时排名','本人当时排名','种子','结果','状态','高芙视角比分','计入实打胜负','赛事最终积分不可逐场相加','WTA比赛ID','feed胜者码','feed比赛状态','官方逐场来源','官方日期来源','旧包TE日期仅供比较']
for title,rows in [('62场实打',[x for x in D['matches']if x['count_wl']]),('74条含轮空弃权',D['matches']),('29场三盘',[x for x in D['matches']if x['reason_code']=='W'and len(x['sets'])==3])]:
 s=wb.create_sheet(title);s.append(heads)
 for m in rows:s.append([m['event'],m['round_display'],m['round'],m['date'],m['date_bjt'],m['date_utc'],m['date_basis'],m['start_utc'],m.get('start_local'),m.get('start_bjt'),None,m['timezone'],cn(m['opponent'])if m['opponent']else'',m['opponent'],m['opponent_rank'],m['player_rank'],m['seed'],m['result'],m['status'],m['score_player'],'是'if m['count_wl']else'否',m['event_final_points'],m.get('match_id'),m.get('winner_code'),m.get('match_state'),m['player_api_source'],m.get('date_source'),m['old_date']])
s=wb.create_sheet('时间线与提纲');s.append(['时段估算','章节','画面原声计划','事实与赛事积分'])
for row in O['rows']:s.append([row['time'],row['title'],row['visual'],row['story']])
s.append(['说明','306秒为第一版估算','所有8场所示比赛均需真实完整MP；额外暂定USQF救MP1个，源绑定待核','不是用户逐字批准，也不是最终成片QC'])
s=wb.create_sheet('退赛与资格');s.append(['事件','时间口径','事实','官方来源'])
for row in [('北京末场','2026-10-07当地','第四轮负梅尔滕斯6-3/6-3（胜者视角）','https://www.wtatennis.com/news/4587818'),('武汉退赛','2026-10-07T17:08:00Z官方报道','Wuhan官方IGStory转引本人：漫长一年后身体需要休息恢复；不是在武汉输球','https://www.wtatennis.com/news/4588144/coco-gauff-naomi-osaka-to-miss-wuhan-open'),('年终正式确认','2026-10-14T10:00:00Z＝北京时间18:00','连续第五年入围；最早数学锁定时刻未证实',D['qualification']['official_source']),('年终尚未开赛','以制作时已知事实为界','不预写年终赛果；曾获2024年终冠军不等于2025卫冕冠军','https://www.wtatennis.com/news/4591678')]:s.append(row)
s=wb.create_sheet('冲突与原始证据');s.append(['类型','状态','原始字段或限制','来源'])
for issue in D['issues']:s.append([issue['kind'],issue['status'],json.dumps(issue,ensure_ascii=False),'https://www.wtatennis.com/news/4504616'])
for tb in D['tiebreak_checks']:s.append(['抢七字段',tb['status'],json.dumps(tb,ensure_ascii=False),'官方playerhistory＋tournamentfeed'])
s=wb.create_sheet('来源与SHA');s.append(['本地原始证据','SHA256','说明'])
for p in [R/'facts/canonical.json',R/'facts/raw/player-matches.json',R/'facts/raw/player-page.html',R/'facts/raw/race-2026-10-19.json',R/'facts/raw/event-retrieval.json',R/'facts/raw/unitedcup-retrieval.json',R/'facts/raw/article-4504616.json']:s.append([str(p),hashlib.sha256(p.read_bytes()).hexdigest(),'真实保存的官方原始内容或结构化审计'])
for sheet in wb:
 sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.zoomScale=85
 for c in sheet[1]:c.font=Font(name='微软雅黑',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17324C');c.alignment=Alignment(wrap_text=True,vertical='center')
 for row in sheet.iter_rows(min_row=2):
  for c in row:c.font=Font(name='微软雅黑',size=10);c.alignment=Alignment(wrap_text=True,vertical='top')
 sheet.row_dimensions[1].height=34
 for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(48,max(14,max(len(str(x.value or''))for x in col)*.9+2))
 if sheet.title=='口径与核验':sheet.column_dimensions['A'].width=26;sheet.column_dimensions['B'].width=110
p=R/'delivery/高芙2026赛季全数据.xlsx';wb.save(p);d=load_workbook(p,read_only=True,data_only=True);assert d['16站汇总'].max_row==17 and d['62场实打'].max_row==63 and d['74条含轮空弃权'].max_row==75 and d['29场三盘'].max_row==30
assert sum(row[11]for row in d['16站汇总'].iter_rows(min_row=2,values_only=True))==5774
(R/'delivery/workbook-validation.json').write_text(json.dumps({'xlsx':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sheets':d.sheetnames,'16events':True,'62played':True,'74entries':True,'29three_set':True,'points5774':True,'rome_tiebreak_conflict_open':True},ensure_ascii=False,indent=2)+'\n');print(p)
